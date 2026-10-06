"""
product_variant_loader.py
--------------------------
Loads and indexes product variant data from the Odoo Product Variant export
Excel file (product.product).

The exported file always has three sheets with slightly different layouts:
  Sheet1  -- has a header row; largest dataset
            Columns: ID | InternalReference | Barcode | Name | Display Name |
                     HSN/SAC Code | Sales Price | Cost
  Sheet2  -- NO header row; same column order as Sheet1 data rows
            (ID at col-0, InternalReference at col-1, Barcode at col-2, ...)
  Sheet3  -- has a header row with Odoo-style prefixed names
            Columns: ID | Product/Barcode | Product/Internal Reference |
                     Product/HSN/SAC Code | Product/Sales Price | Product/Cost
            (plus extra cluster/store/transfer columns that are ignored here)

Lookup order: Sheet1 -> Sheet2 -> Sheet3 (first match wins).
Products can be found by barcode OR by internal reference (SKU).
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

@dataclass
class ProductVariant:
    """Normalised product variant record built from any of the three sheets."""
    odoo_id: str
    internal_ref: str
    barcode: str
    name: str
    display_name: str
    hsn_code: str
    sales_price: float
    cost: float
    source_sheet: str

    def __repr__(self) -> str:
        return (
            f"ProductVariant(sku={self.internal_ref!r}, barcode={self.barcode!r}, "
            f"name={self.name!r}, sheet={self.source_sheet!r})"
        )


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_barcode(value: Any) -> str:
    raw = _clean(value)
    if not raw:
        return ""
    if raw.endswith(".0") and raw[:-2].isdigit():
        raw = raw[:-2]
    return raw


def _normalize_sku(value: Any) -> str:
    return _clean(value).upper()


# ---------------------------------------------------------------------------
# Per-sheet parsers
# ---------------------------------------------------------------------------

def _parse_sheet1(ws) -> list[ProductVariant]:
    records: list[ProductVariant] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        odoo_id      = _clean(row[0])
        internal_ref = _clean(row[1])
        barcode      = _normalize_barcode(row[2])
        name         = _clean(row[3])
        display_name = _clean(row[4])
        hsn_code     = _clean(row[5])
        sales_price  = _to_float(row[6])
        cost         = _to_float(row[7])
        if not internal_ref and not barcode:
            continue
        records.append(ProductVariant(
            odoo_id=odoo_id, internal_ref=internal_ref, barcode=barcode,
            name=name, display_name=display_name, hsn_code=hsn_code,
            sales_price=sales_price, cost=cost, source_sheet="Sheet1",
        ))
    return records


def _parse_sheet2(ws) -> list[ProductVariant]:
    records: list[ProductVariant] = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        odoo_id      = _clean(row[0])
        internal_ref = _clean(row[1])
        barcode      = _normalize_barcode(row[2])
        name         = _clean(row[3])
        display_name = _clean(row[4])
        hsn_code     = _clean(row[5])
        sales_price  = _to_float(row[6])
        cost         = _to_float(row[7])
        if not internal_ref and not barcode:
            continue
        records.append(ProductVariant(
            odoo_id=odoo_id, internal_ref=internal_ref, barcode=barcode,
            name=name, display_name=display_name, hsn_code=hsn_code,
            sales_price=sales_price, cost=cost, source_sheet="Sheet2",
        ))
    return records


def _parse_sheet3(ws) -> list[ProductVariant]:
    records: list[ProductVariant] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        odoo_id      = _clean(row[0])
        barcode      = _normalize_barcode(row[1])
        internal_ref = _clean(row[2])
        hsn_code     = _clean(row[3])
        sales_price  = _to_float(row[4])
        cost         = _to_float(row[5])
        name         = ""
        display_name = ""
        if not internal_ref and not barcode:
            continue
        records.append(ProductVariant(
            odoo_id=odoo_id, internal_ref=internal_ref, barcode=barcode,
            name=name, display_name=display_name, hsn_code=hsn_code,
            sales_price=sales_price, cost=cost, source_sheet="Sheet3",
        ))
    return records


# ---------------------------------------------------------------------------
# Public API: ProductVariantIndex
# ---------------------------------------------------------------------------

@dataclass
class ProductVariantIndex:
    """
    In-memory index of all product variants loaded from the three-sheet
    Odoo export.  Provides O(1) lookup by barcode or internal reference.

    Lookup order: Sheet1 -> Sheet2 -> Sheet3  (first match wins).
    """
    _by_barcode: dict[str, ProductVariant] = field(default_factory=dict, repr=False)
    _by_sku: dict[str, ProductVariant] = field(default_factory=dict, repr=False)
    _stats: dict[str, int] = field(default_factory=dict, repr=False)

    @classmethod
    def load(cls, xlsx_path) -> "ProductVariantIndex":
        path = Path(xlsx_path)
        if not path.exists():
            raise FileNotFoundError(f"Product variant file not found: {path}")

        _log.info("Loading product variants from %s ...", path.name)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet_names = wb.sheetnames
            all_records = []

            if "Sheet1" in sheet_names:
                s1 = _parse_sheet1(wb["Sheet1"])
                _log.info("Sheet1: %d records", len(s1))
                all_records.extend(s1)
            else:
                _log.warning("Sheet1 not found in %s", path.name)

            if "Sheet2" in sheet_names:
                s2 = _parse_sheet2(wb["Sheet2"])
                _log.info("Sheet2: %d records", len(s2))
                all_records.extend(s2)
            else:
                _log.warning("Sheet2 not found in %s", path.name)

            if "Sheet3" in sheet_names:
                s3 = _parse_sheet3(wb["Sheet3"])
                _log.info("Sheet3: %d records", len(s3))
                all_records.extend(s3)
            else:
                _log.warning("Sheet3 not found in %s", path.name)
        finally:
            wb.close()

        by_barcode = {}
        by_sku = {}
        dup_barcode = dup_sku = 0

        for rec in all_records:
            bc_key  = rec.barcode.upper()      if rec.barcode      else ""
            sku_key = rec.internal_ref.upper() if rec.internal_ref else ""

            if bc_key and bc_key not in by_barcode:
                by_barcode[bc_key] = rec
            elif bc_key:
                dup_barcode += 1

            if sku_key and sku_key not in by_sku:
                by_sku[sku_key] = rec
            elif sku_key:
                dup_sku += 1

        stats = {
            "total_records": len(all_records),
            "unique_barcodes": len(by_barcode),
            "unique_skus": len(by_sku),
            "duplicate_barcodes_skipped": dup_barcode,
            "duplicate_skus_skipped": dup_sku,
        }
        _log.info(
            "ProductVariantIndex ready: %d total, %d unique barcodes, %d unique SKUs",
            stats["total_records"], stats["unique_barcodes"], stats["unique_skus"],
        )

        idx = cls()
        idx._by_barcode = by_barcode
        idx._by_sku = by_sku
        idx._stats = stats
        return idx

    def find_by_barcode(self, barcode) -> "ProductVariant | None":
        key = _normalize_barcode(barcode).upper()
        return self._by_barcode.get(key)

    def find_by_sku(self, sku) -> "ProductVariant | None":
        key = _normalize_sku(sku).upper()
        return self._by_sku.get(key)

    def find(self, barcode="", sku="") -> "ProductVariant | None":
        """Find by barcode first; fall back to SKU. Sheet1 -> Sheet2 -> Sheet3."""
        if barcode:
            result = self.find_by_barcode(barcode)
            if result is not None:
                return result
        if sku:
            result = self.find_by_sku(sku)
            if result is not None:
                return result
        return None

    @property
    def stats(self) -> dict:
        return dict(self._stats)

    def __len__(self) -> int:
        return self._stats.get("total_records", 0)


# ---------------------------------------------------------------------------
# Module-level singleton (lazy-loaded, cached)
# ---------------------------------------------------------------------------

_CACHED_INDEX = None
_CACHED_PATH  = ""


def get_product_variant_index(xlsx_path) -> ProductVariantIndex:
    """
    Return a cached ProductVariantIndex for the given file.
    Reloads automatically if a different path is given.

    Usage::

        from product_variant_loader import get_product_variant_index
        index = get_product_variant_index("path/to/Product Variant (...).xlsx")
        variant = index.find(barcode="8905639008385", sku="BSB104NVY032")
        if variant:
            print(variant.hsn_code, variant.sales_price)
        else:
            print("Not found in any sheet")
    """
    global _CACHED_INDEX, _CACHED_PATH
    path_str = str(Path(xlsx_path).resolve())
    if _CACHED_INDEX is None or _CACHED_PATH != path_str:
        _CACHED_INDEX = ProductVariantIndex.load(xlsx_path)
        _CACHED_PATH  = path_str
    return _CACHED_INDEX


# ---------------------------------------------------------------------------
# Quick self-test (run this file directly to verify)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    import sys, argparse
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    ap = argparse.ArgumentParser(description="Test product variant lookup across 3 sheets")
    ap.add_argument("xlsx", help="Path to the Product Variant Excel file")
    ap.add_argument("--barcode", help="Barcode to look up")
    ap.add_argument("--sku",     help="Internal reference (SKU) to look up")
    args = ap.parse_args()

    index = ProductVariantIndex.load(args.xlsx)
    print("\n=== Stats ===")
    for k, v in index.stats.items():
        print(f"  {k}: {v}")

    if args.barcode or args.sku:
        print(f"\n=== Lookup: barcode={args.barcode!r}  sku={args.sku!r} ===")
        result = index.find(barcode=args.barcode or "", sku=args.sku or "")
        if result:
            print(f"  FOUND in {result.source_sheet}")
            print(f"  ID          : {result.odoo_id}")
            print(f"  SKU         : {result.internal_ref}")
            print(f"  Barcode     : {result.barcode}")
            print(f"  Name        : {result.name}")
            print(f"  Display Name: {result.display_name}")
            print(f"  HSN         : {result.hsn_code}")
            print(f"  Sales Price : {result.sales_price}")
            print(f"  Cost        : {result.cost}")
        else:
            print("  NOT FOUND in any sheet (Sheet1, Sheet2, Sheet3)")
