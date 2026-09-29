from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="312E81")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUCCESS_FILL = PatternFill("solid", fgColor="DCFCE7")
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")
WARN_FILL = PatternFill("solid", fgColor="FEF3C7")


def _style_sheet(ws):
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for row in range(1, min(ws.max_row, 300) + 1):
            value = ws.cell(row, col).value
            max_len = max(max_len, len(str(value)) if value is not None else 0)
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 11), 42)


def write_result(path: Path, documents: list[dict], post_results: list[dict] | None = None, api_calls: list[dict] | None = None, validation_status: str = "VALIDATED"):
    result_map = {r["document_key"]: r for r in (post_results or [])}
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    headers = [
        "Source Type", "Odoo Reference", "Document Date", "Store", "Destination Site Code", "State", "Transfer Doc Code",
        "Item Lines", "Total Qty", "DC Code", "DC Number", "Transfer Out Code", "Transfer Out Number", "Status", "Portal User",
        "DC Ginesys User", "DC Ginesys User Code", "Transfer Out Ginesys User", "Transfer Out Ginesys User Code", "Error"
    ]
    ws.append(headers)
    for doc in documents:
        r = result_map.get(doc["document_key"], {})
        status = r.get("status") or validation_status
        ws.append([
            doc["source_type"], doc["reference"], doc["date"], doc["store_name"], doc["site_code"], doc.get("state_code"),
            doc["invoice_doc_code"], doc["line_count"], doc["total_qty"], r.get("dc_code"), r.get("dc_number"),
            r.get("transfer_code"), r.get("transfer_number"), status, doc.get("posted_by_display") or doc.get("posted_by_username"),
            r.get("dc_created_by_ginesys_username"), r.get("dc_created_by_ginesys_user_code"),
            r.get("transfer_created_by_ginesys_username"), r.get("transfer_created_by_ginesys_user_code"), r.get("error", "")
        ])
        fill = SUCCESS_FILL if status in {"SUCCESS", "ALREADY_POSTED"} else FAIL_FILL if "FAIL" in status else WARN_FILL if status == "BLOCKED" else None
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill
    _style_sheet(ws)

    items = wb.create_sheet("Items")
    items.append([
        "Odoo Reference", "Store", "Site Code", "Odoo SKU", "Ginesys Item", "Barcode", "Qty",
        "MRP", "Basic Rate", "Discount Factor", "Discount", "Final Rate", "WSP", "Pricing Source"
    ])
    for doc in documents:
        for item in doc["items"]:
            items.append([
                doc["reference"], doc["store_name"], doc["site_code"], item["sku"], item.get("ginesys_item_id"),
                item.get("barcode"), item["qty"], item.get("mrp"), item.get("basic_rate"),
                item.get("discount_factor", item.get("factor")), item.get("discount"), item.get("rate"), item.get("wsp"),
                "Utility/SelectItem exact barcode lookup"
            ])
    _style_sheet(items)

    if api_calls is not None:
        logs = wb.create_sheet("API_Log")
        logs.append(["Timestamp", "Ginesys User", "Endpoint", "HTTP Status", "Success", "Message", "Request ID", "Duration ms"])
        now = datetime.now().isoformat(timespec="seconds")
        for call in api_calls:
            logs.append([now, call.get("ginesys_username"), call.get("endpoint"), call.get("http_status"), call.get("success"), call.get("message"), call.get("request_id"), call.get("duration_ms")])
        _style_sheet(logs)

    wb.save(path)
